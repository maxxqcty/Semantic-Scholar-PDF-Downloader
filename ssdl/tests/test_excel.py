from openpyxl import load_workbook

from ssdl.excel import create_excel

HEADERS = [
    "Number", "Title", "Authors", "Year", "Venue",
    "DOI", "Paper URL", "PDF URL", "SHA256", "Filename",
]

RECORD = {
    "number": 1,
    "title": "Deep Learning",
    "authors": "A. Author, B. Author",
    "year": 2024,
    "venue": "NeurIPS",
    "doi": "10.1/xyz",
    "paper_url": "https://www.semanticscholar.org/paper/abc",
    "pdf_url": "https://example.org/abc.pdf",
    "sha256": "deadbeef",
    "filename": "01 - Deep Learning [abc].pdf",
}


def test_writes_headers_and_record_row(tmp_path):
    path = tmp_path / "papers.xlsx"

    create_excel(path, [RECORD])

    ws = load_workbook(path).active
    assert ws.title == "Downloaded Papers"
    assert [ws.cell(row=1, column=i + 1).value for i in range(10)] == HEADERS
    assert [ws.cell(row=2, column=i + 1).value for i in range(10)] == list(RECORD.values())


def test_empty_records_writes_header_only(tmp_path):
    path = tmp_path / "papers.xlsx"

    create_excel(path, [])

    ws = load_workbook(path).active
    assert ws.max_row == 1
    assert ws.cell(row=1, column=1).value == "Number"


def test_freezes_header_and_links_urls(tmp_path):
    path = tmp_path / "papers.xlsx"

    create_excel(path, [RECORD])

    ws = load_workbook(path).active
    assert ws.freeze_panes == "A2"
    assert ws.cell(row=2, column=7).hyperlink.target == RECORD["paper_url"]
    assert ws.cell(row=2, column=8).hyperlink.target == RECORD["pdf_url"]


def test_empty_urls_get_no_hyperlink(tmp_path):
    path = tmp_path / "papers.xlsx"
    record = dict(RECORD, paper_url="", pdf_url="")

    create_excel(path, [record])

    ws = load_workbook(path).active
    assert ws.cell(row=2, column=7).hyperlink is None
    assert ws.cell(row=2, column=8).hyperlink is None


def test_overwrites_previous_workbook(tmp_path):
    path = tmp_path / "papers.xlsx"
    create_excel(path, [RECORD])

    create_excel(path, [])

    ws = load_workbook(path).active
    assert ws.max_row == 1
