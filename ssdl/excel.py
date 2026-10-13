
from openpyxl import Workbook


def create_excel(
    excel_path,
    records,
):
    """
    Create an Excel workbook containing only
    successfully downloaded PDFs.
    """

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Downloaded Papers"

    headers = [
        "Number",
        "Title",
        "Authors",
        "Year",
        "Venue",
        "DOI",
        "Paper URL",
        "PDF URL",
        "SHA256",
        "Filename",
    ]

    worksheet.append(headers)

    for record in records:

        worksheet.append([
            record["number"],
            record["title"],
            record["authors"],
            record["year"],
            record["venue"],
            record["doi"],
            record["paper_url"],
            record["pdf_url"],
            record["sha256"],
            record["filename"],
        ])

    # Freeze the header row.
    worksheet.freeze_panes = "A2"

    # Make URLs clickable.
    for row in range(
        2,
        worksheet.max_row + 1,
    ):

        paper_url_cell = worksheet.cell(
            row=row,
            column=7,
        )

        pdf_url_cell = worksheet.cell(
            row=row,
            column=8,
        )

        if paper_url_cell.value:
            paper_url_cell.hyperlink = (
                paper_url_cell.value
            )

        if pdf_url_cell.value:
            pdf_url_cell.hyperlink = (
                pdf_url_cell.value
            )

    # Set column widths.
    widths = {
        "A": 10,
        "B": 60,
        "C": 45,
        "D": 10,
        "E": 30,
        "F": 35,
        "G": 65,
        "H": 65,
        "I": 70,
        "J": 70,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[
            column
        ].width = width

    workbook.save(excel_path)
