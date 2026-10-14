import hashlib

import pytest
from openpyxl import load_workbook

from ssdl import cli


def test_help_exits_zero(capsys):
    with pytest.raises(SystemExit) as exc:
        cli.main(["--help"])
    assert exc.value.code == 0
    output = capsys.readouterr().out
    assert "usage:" in output
    assert "--query" in output


def test_missing_query_is_rejected():
    with pytest.raises(SystemExit) as exc:
        cli.main([])
    assert exc.value.code == 2


@pytest.mark.parametrize("extra", [["--max-results", "0"], ["--page-size", "0"],
                                   ["--page-size", "101"], ["--delay", "-1"],
                                   ["--search-delay", "-1"]])
def test_invalid_arguments_are_rejected(extra):
    with pytest.raises(SystemExit) as exc:
        cli.main(["--query", "q"] + extra)
    assert exc.value.code == 2


def _stub(monkeypatch, page):
    monkeypatch.setattr(cli, "create_session", lambda: object())
    monkeypatch.setattr(cli, "search_papers", lambda **kwargs: page)


def _fake_download(url, destination):
    destination.write_bytes(b"%PDF-1.4 fake")
    return True, "Downloaded"


def test_full_run_writes_pdf_and_excel(tmp_path, monkeypatch):
    _stub(monkeypatch, {
        "data": [
            {"paperId": "aaa11122233344455", "title": "No Open Access"},
            {
                "paperId": "abcdef1234567890",
                "title": "Deep Learning",
                "authors": [{"name": "A. Author"}],
                "year": 2024,
                "venue": "NeurIPS",
                "externalIds": {"DOI": "10.1/xyz"},
                "url": "https://www.semanticscholar.org/paper/abcdef",
                "openAccessPdf": {"url": "https://example.org/paper.pdf"},
            },
        ],
        "total": 2,
        "next": None,
    })
    monkeypatch.setattr(
        cli, "download_pdf",
        lambda url, destination, session, max_retries=4: _fake_download(url, destination),
    )

    out = tmp_path / "out"
    cli.main(["--query", "q", "--max-results", "1", "--output", str(out),
              "--delay", "0", "--search-delay", "0"])

    assert [p.name for p in out.glob("*.pdf")] == ["01 - Deep Learning [abcdef123456].pdf"]

    ws = load_workbook(out / "papers.xlsx").active
    assert ws.max_row == 2
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=2).value == "Deep Learning"
    assert ws.cell(row=2, column=9).value == hashlib.sha256(b"%PDF-1.4 fake").hexdigest()
    assert ws.cell(row=2, column=10).value == "01 - Deep Learning [abcdef123456].pdf"


def test_empty_prefix_is_not_inserted_into_filename(tmp_path, monkeypatch):
    _stub(monkeypatch, {
        "data": [{"paperId": "111111111111aaaa", "title": "First",
                  "openAccessPdf": {"url": "https://example.org/1.pdf"}}],
        "total": 1,
        "next": None,
    })
    monkeypatch.setattr(
        cli, "download_pdf",
        lambda url, destination, session, max_retries=4: _fake_download(url, destination),
    )

    out = tmp_path / "out"
    cli.main(["--query", "q", "--max-results", "1", "--prefix", "", "--output", str(out),
              "--delay", "0", "--search-delay", "0"])

    assert [p.name for p in out.glob("*.pdf")] == ["01 - First [111111111111].pdf"]


def test_duplicate_content_is_skipped_and_numbering_continues(tmp_path, monkeypatch):
    def paper(pid, title):
        return {"paperId": pid, "title": title,
                "openAccessPdf": {"url": f"https://example.org/{pid}.pdf"}}

    _stub(monkeypatch, {
        "data": [paper("111111111111aaaa", "First"),
                 paper("222222222222bbbb", "Second"),
                 paper("333333333333cccc", "Third")],
        "total": 3,
        "next": None,
    })
    bodies = [b"%PDF-1.4 same", b"%PDF-1.4 same", b"%PDF-1.4 other"]

    def fake_download(url, destination, session, max_retries=4):
        destination.write_bytes(bodies.pop(0))
        return True, "Downloaded"

    monkeypatch.setattr(cli, "download_pdf", fake_download)

    out = tmp_path / "out"
    cli.main(["--query", "q", "--max-results", "2", "--output", str(out),
              "--delay", "0", "--search-delay", "0"])

    names = sorted(p.name for p in out.glob("*.pdf"))
    assert names == ["01 - First [111111111111].pdf",
                     "02 - Third [333333333333].pdf"]

    ws = load_workbook(out / "papers.xlsx").active
    assert ws.max_row == 3
    assert ws.cell(row=3, column=1).value == 2
    assert ws.cell(row=3, column=2).value == "Third"


def test_default_output_anchors_to_project_root(tmp_path, monkeypatch):
    fake_root = tmp_path / "proj"
    fake_pkg = fake_root / "ssdl"
    fake_pkg.mkdir(parents=True)
    monkeypatch.setattr(cli, "__file__", str(fake_pkg / "cli.py"))
    _stub(monkeypatch, {"data": [], "total": 0, "next": None})

    cli.main(["--query", "q", "--delay", "0", "--search-delay", "0"])

    assert (fake_root / "papers" / "papers.xlsx").is_file()
    assert list(fake_pkg.iterdir()) == []


def test_second_page_is_fetched_when_target_exceeds_first_page(tmp_path, monkeypatch):
    """Crossing a page boundary must advance the offset and keep going."""

    def paper(pid, title, with_pdf):
        entry = {"paperId": pid, "title": title}
        if with_pdf:
            entry["openAccessPdf"] = {"url": f"https://example.org/{pid}.pdf"}
        return entry

    page_one = {"data": [paper(f"000000000000{i:04d}", f"Page One {i}", False)
                         for i in range(100)],
                "total": 101, "next": 100}
    page_two = {"data": [paper("999999999999xxxx", "Page Two", True)],
                "total": 101, "next": None}

    monkeypatch.setattr(cli, "create_session", lambda: object())
    calls = []

    def fake_search(query, offset, limit, session):
        # Requesting the same offset twice is the pagination
        # failure mode that would otherwise spin forever.
        if offset in calls:
            raise RuntimeError(f"pagination stuck at offset {offset}")
        calls.append(offset)
        return page_one if offset == 0 else page_two

    monkeypatch.setattr(cli, "search_papers", fake_search)
    monkeypatch.setattr(
        cli, "download_pdf",
        lambda url, destination, session, max_retries=4: _fake_download(url, destination),
    )

    out = tmp_path / "out"
    cli.main(["--query", "q", "--max-results", "1", "--output", str(out),
              "--delay", "0", "--search-delay", "0"])

    assert calls == [0, 100]
    assert [p.name for p in out.glob("*.pdf")] == ["01 - Page Two [999999999999].pdf"]


def test_absolute_output_is_used_as_is(tmp_path, monkeypatch):
    _stub(monkeypatch, {"data": [], "total": 0, "next": None})
    out = tmp_path / "elsewhere"

    cli.main(["--query", "q", "--output", str(out), "--delay", "0", "--search-delay", "0"])

    assert (out / "papers.xlsx").is_file()
