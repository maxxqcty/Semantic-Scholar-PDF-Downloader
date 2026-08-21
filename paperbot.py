
#!/usr/bin/env python3

import argparse
import hashlib
import re
import time
from pathlib import Path

import requests
from openpyxl import Workbook


API_URL = "https://api.semanticscholar.org/graph/v1/paper/search"

FIELDS = ",".join([
    "paperId",
    "title",
    "authors",
    "year",
    "venue",
    "url",
    "externalIds",
    "openAccessPdf",
])


def safe_filename(text, max_length=150):
    """
    Convert text into a filename-safe string.
    """

    text = text or "Untitled"

    # Characters forbidden in Windows filenames:
    # < > : " / \ | ? *
    text = re.sub(r'[<>:"/\\|?*]', "", text)

    # Normalize whitespace.
    text = re.sub(r"\s+", " ", text).strip()

    # Avoid filenames ending in a period or space.
    return text[:max_length].rstrip(". ")


def file_hash(path):
    """
    Calculate SHA-256 hash of a file.
    """

    sha256 = hashlib.sha256()

    with open(path, "rb") as file:
        for chunk in iter(
            lambda: file.read(1024 * 1024),
            b"",
        ):
            sha256.update(chunk)

    return sha256.hexdigest()


def create_session():
    """
    Create the HTTP session.

    No API key is required.
    """

    session = requests.Session()

    session.headers.update({
        "User-Agent": (
            "ResearchPDFDownloader/3.0 "
            "(academic research)"
        ),
        "Accept": "application/json",
    })

    return session


def search_papers(
    query,
    offset,
    limit,
    session,
    max_retries=7,
):
    """
    Search Semantic Scholar.

    Handles:
        429 rate limiting
        500 server errors
        502 bad gateway
        503 service unavailable
        504 gateway timeout
    """

    params = {
        "query": query,
        "offset": offset,
        "limit": min(limit, 100),
        "fields": FIELDS,

        # Request papers for which Semantic Scholar
        # reports an open-access PDF.
        "openAccessPdf": "",
    }

    for attempt in range(max_retries):

        try:
            response = session.get(
                API_URL,
                params=params,
                timeout=30,
            )

        except requests.RequestException as exc:

            if attempt == max_retries - 1:
                raise

            wait = min(
                120,
                2 ** attempt,
            )

            print(
                f"  Search connection error: {exc}"
            )

            print(
                f"  Retrying in {wait:.0f} seconds..."
            )

            time.sleep(wait)
            continue

        # Successful request.
        if response.status_code == 200:
            return response.json()

        # Rate limited.
        if response.status_code == 429:

            retry_after = response.headers.get(
                "Retry-After"
            )

            if retry_after:
                try:
                    wait = float(retry_after)
                except ValueError:
                    wait = min(
                        120,
                        5 * (2 ** attempt),
                    )
            else:
                wait = min(
                    120,
                    5 * (2 ** attempt),
                )

            print()
            print(
                "  Semantic Scholar rate limit "
                "reached (HTTP 429)."
            )

            print(
                f"  Waiting {wait:.0f} seconds "
                "before retrying..."
            )

            time.sleep(wait)
            continue

        # Temporary server errors.
        if response.status_code in {
            500,
            502,
            503,
            504,
        }:

            if attempt == max_retries - 1:
                response.raise_for_status()

            wait = min(
                60,
                2 ** attempt,
            )

            print(
                f"  Semantic Scholar server error "
                f"({response.status_code})."
            )

            print(
                f"  Retrying in {wait:.0f} seconds..."
            )

            time.sleep(wait)
            continue

        # Other HTTP errors.
        response.raise_for_status()

    raise RuntimeError(
        "Semantic Scholar search failed after "
        f"{max_retries} attempts."
    )


def download_pdf(
    url,
    destination,
    session,
    max_retries=4,
):
    """
    Download and validate a PDF.

    Uses a .part file so interrupted downloads
    do not leave incomplete PDFs.
    """

    temporary = destination.with_suffix(
        destination.suffix + ".part"
    )

    for attempt in range(max_retries):

        try:

            with session.get(
                url,
                timeout=60,
                stream=True,
                allow_redirects=True,
            ) as response:

                # Handle rate limits and temporary
                # server errors.
                if response.status_code in {
                    429,
                    500,
                    502,
                    503,
                    504,
                }:

                    retry_after = response.headers.get(
                        "Retry-After"
                    )

                    if retry_after:
                        try:
                            wait = float(retry_after)
                        except ValueError:
                            wait = min(
                                60,
                                3 * (2 ** attempt),
                            )
                    else:
                        wait = min(
                            60,
                            3 * (2 ** attempt),
                        )

                    print(
                        f"  PDF server returned "
                        f"HTTP {response.status_code}."
                    )

                    print(
                        f"  Retrying in "
                        f"{wait:.0f} seconds..."
                    )

                    time.sleep(wait)
                    continue

                response.raise_for_status()

                content_type = response.headers.get(
                    "content-type",
                    "",
                ).lower()

                first_chunk = next(
                    response.iter_content(
                        chunk_size=8192
                    ),
                    b"",
                )

                if not first_chunk:
                    return False, "Empty response"

                # Check both content type and PDF signature.
                is_pdf = (
                    "application/pdf" in content_type
                    or first_chunk.startswith(b"%PDF")
                )

                if not is_pdf:
                    return False, (
                        "Response is not a PDF "
                        f"(content-type: "
                        f"{content_type or 'unknown'})"
                    )

                with open(
                    temporary,
                    "wb",
                ) as file:

                    file.write(first_chunk)

                    for chunk in response.iter_content(
                        chunk_size=1024 * 1024
                    ):
                        if chunk:
                            file.write(chunk)

            # Download completed.
            temporary.replace(destination)

            return True, "Downloaded"

        except requests.RequestException as exc:

            temporary.unlink(
                missing_ok=True
            )

            if attempt == max_retries - 1:
                return False, str(exc)

            wait = min(
                60,
                2 ** attempt,
            )

            print(
                f"  Download error: {exc}"
            )

            print(
                f"  Retrying in "
                f"{wait:.0f} seconds..."
            )

            time.sleep(wait)

        except OSError as exc:

            temporary.unlink(
                missing_ok=True
            )

            return False, str(exc)

    return False, "Download failed"


def create_excel(
    excel_path,
    records,
):
    """
    Create an Excel workbook containing only
    successfully downloaded PDFs.
    """

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Downloaded Papers"

    headers = [
        "Number",
        "Title",
        "Authors",
        "Year",
        "Venue",
        "DOI",
        "Paper URL",
        "PDF URL",
        "SHA256",
        "Filename",
    ]

    worksheet.append(headers)

    for record in records:

        worksheet.append([
            record["number"],
            record["title"],
            record["authors"],
            record["year"],
            record["venue"],
            record["doi"],
            record["paper_url"],
            record["pdf_url"],
            record["sha256"],
            record["filename"],
        ])

    # Freeze the header row.
    worksheet.freeze_panes = "A2"

    # Make URLs clickable.
    for row in range(
        2,
        worksheet.max_row + 1,
    ):

        paper_url_cell = worksheet.cell(
            row=row,
            column=7,
        )

        pdf_url_cell = worksheet.cell(
            row=row,
            column=8,
        )

        if paper_url_cell.value:
            paper_url_cell.hyperlink = (
                paper_url_cell.value
            )

        if pdf_url_cell.value:
            pdf_url_cell.hyperlink = (
                pdf_url_cell.value
            )

    # Set column widths.
    widths = {
        "A": 10,
        "B": 60,
        "C": 45,
        "D": 10,
        "E": 30,
        "F": 35,
        "G": 65,
        "H": 65,
        "I": 70,
        "J": 70,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[
            column
        ].width = width

    workbook.save(excel_path)


def main():

    parser = argparse.ArgumentParser(
        description=(
            "Download open-access research PDFs "
            "from Semantic Scholar."
        )
    )

    parser.add_argument(
        "--query",
        required=True,
        help="Research query.",
    )

    parser.add_argument(
        "--max-results",
        type=int,
        default=50,
        help=(
            "Number of successfully downloaded "
            "unique PDFs. Default: 50."
        ),
    )

    parser.add_argument(
        "--output",
        default="papers",
        help=(
            "Download directory. "
            "Default: papers"
        ),
    )

    parser.add_argument(
        "--prefix",
        default="",
        help=(
            "Optional prefix for PDF filenames. "
            "Example: --prefix 'ML Healthcare'"
        ),
    )

    parser.add_argument(
        "--page-size",
        type=int,
        default=100,
        help=(
            "Number of search results per page. "
            "Maximum: 100. Default: 100."
        ),
    )

    parser.add_argument(
        "--delay",
        type=float,
        default=2.0,
        help=(
            "Seconds between PDF downloads. "
            "Default: 2."
        ),
    )

    parser.add_argument(
        "--search-delay",
        type=float,
        default=3.0,
        help=(
            "Seconds between search pages. "
            "Default: 3."
        ),
    )

    args = parser.parse_args()

    if args.max_results <= 0:
        parser.error(
            "--max-results must be greater than 0"
        )

    if not 1 <= args.page_size <= 100:
        parser.error(
            "--page-size must be between 1 and 100"
        )

    if args.delay < 0:
        parser.error(
            "--delay cannot be negative"
        )

    if args.search_delay < 0:
        parser.error(
            "--search-delay cannot be negative"
        )

    # Convert the output directory to an absolute path.
    script_dir = Path(__file__).resolve().parent

    output_path = Path(args.output)

    if output_path.is_absolute():
        output_dir = output_path
    else:
        output_dir = script_dir / output_path

    output_dir = output_dir.resolve()

    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    excel_path = (
        output_dir / "papers.xlsx"
    )

    session = create_session()

    # Clean the optional filename prefix.
    filename_prefix = safe_filename(
        args.prefix
    )

    print("=" * 70)
    print(
        "Semantic Scholar Open-Access PDF Downloader"
    )
    print("=" * 70)
    print()

    print(
        f"Query:           {args.query}"
    )

    print(
        f"Target PDFs:     {args.max_results}"
    )

    print(
        f"Download folder: {output_dir}"
    )

    print(
        f"Excel file:      {excel_path}"
    )

    print(
        f"Filename prefix: "
        f"{filename_prefix or '(none)'}"
    )

    print(
        f"Page size:       {args.page_size}"
    )

    print(
        f"PDF delay:       {args.delay}s"
    )

    print(
        f"Search delay:    {args.search_delay}s"
    )

    print()

    print(
        "API key:         Not required"
    )

    print()

    downloaded_records = []

    # Hashes of successfully downloaded PDFs.
    seen_hashes = set()

    # Semantic Scholar paper IDs already processed.
    seen_paper_ids = set()

    offset = 0
    total_available = None

    while (
        len(downloaded_records)
        < args.max_results
    ):

        print("-" * 70)

        print(
            f"Searching results "
            f"{offset + 1}-"
            f"{offset + args.page_size}..."
        )

        try:

            result = search_papers(
                query=args.query,
                offset=offset,
                limit=args.page_size,
                session=session,
            )

        except requests.RequestException as exc:

            print()
            print(
                f"Search failed permanently: {exc}"
            )
            print()
            break

        except Exception as exc:

            print()
            print(
                f"Unexpected search error: {exc}"
            )
            print()
            break

        papers = result.get(
            "data",
            [],
        )

        total_available = result.get(
            "total"
        )

        next_offset = result.get(
            "next"
        )

        print(
            f"Received {len(papers)} results."
        )

        if total_available is not None:

            print(
                "Matching results reported by "
                f"Semantic Scholar: "
                f"{total_available:,}"
            )

        if not papers:

            print(
                "No more search results."
            )

            break

        for paper in papers:

            if (
                len(downloaded_records)
                >= args.max_results
            ):
                break

            paper_id = paper.get(
                "paperId"
            )

            # Avoid processing the same paper twice.
            if (
                paper_id
                and paper_id in seen_paper_ids
            ):
                continue

            if paper_id:
                seen_paper_ids.add(
                    paper_id
                )

            title = (
                paper.get("title")
                or "Untitled"
            )

            authors = ", ".join(
                author.get("name", "")
                for author in (
                    paper.get("authors")
                    or []
                )
                if author.get("name")
            )

            year = (
                paper.get("year")
                or ""
            )

            venue = (
                paper.get("venue")
                or ""
            )

            external_ids = (
                paper.get("externalIds")
                or {}
            )

            doi = external_ids.get(
                "DOI",
                "",
            )

            paper_url = paper.get(
                "url",
                "",
            )

            open_access = (
                paper.get("openAccessPdf")
                or {}
            )

            pdf_url = open_access.get(
                "url"
            )

            print()
            print(
                f"Current successful downloads: "
                f"{len(downloaded_records)}/"
                f"{args.max_results}"
            )

            print(
                f"Paper: {title}"
            )

            if not pdf_url:

                print(
                    "  No PDF URL. Skipping."
                )

                continue

            # The number is based only on successful
            # unique downloads.
            download_number = (
                len(downloaded_records) + 1
            )

            safe_title = safe_filename(
                title
            )

            # Build filename.
            #
            # Without prefix:
            # 01 - Paper Title [paperid].pdf
            #
            # With prefix:
            # 01 - ML Healthcare - Paper Title [paperid].pdf

            filename_parts = [
                f"{download_number:02d}",
            ]

            if filename_prefix:
                filename_parts.append(
                    filename_prefix
                )

            filename_parts.append(
                safe_title
            )

            filename = " - ".join(
                filename_parts
            )

            if paper_id:
                filename += (
                    f" [{paper_id[:12]}]"
                )

            filename += ".pdf"

            destination = (
                output_dir / filename
            )

            # Prevent overwriting an existing file.
            if destination.exists():

                counter = 2

                while True:

                    base_name = (
                        destination.stem
                    )

                    candidate = (
                        output_dir
                        / (
                            f"{base_name} "
                            f"({counter}).pdf"
                        )
                    )

                    if not candidate.exists():

                        destination = candidate
                        break

                    counter += 1

            print(
                f"  PDF URL: {pdf_url}"
            )

            success, message = download_pdf(
                url=pdf_url,
                destination=destination,
                session=session,
            )

            if not success:

                destination.unlink(
                    missing_ok=True
                )

                print(
                    f"  FAILED: {message}"
                )

                continue

            # Calculate SHA-256.
            try:

                digest = file_hash(
                    destination
                )

            except OSError as exc:

                destination.unlink(
                    missing_ok=True
                )

                print(
                    f"  Hashing failed: {exc}"
                )

                continue

            # Detect duplicate PDF content.
            if digest in seen_hashes:

                destination.unlink(
                    missing_ok=True
                )

                print(
                    "  DUPLICATE: PDF deleted."
                )

                continue

            seen_hashes.add(
                digest
            )

            # The PDF is now officially considered
            # a successful download.
            record = {
                "number": download_number,
                "title": title,
                "authors": authors,
                "year": year,
                "venue": venue,
                "doi": doi,
                "paper_url": paper_url,
                "pdf_url": pdf_url,
                "sha256": digest,
                "filename": destination.name,
            }

            downloaded_records.append(
                record
            )

            # Save Excel immediately.
            create_excel(
                excel_path,
                downloaded_records,
            )

            print()
            print(
                "  SUCCESS!"
            )

            print(
                f"  Saved: {destination.name}"
            )

            print(
                f"  Progress: "
                f"{len(downloaded_records)}/"
                f"{args.max_results}"
            )

            print(
                f"  Directory: {output_dir}"
            )

            if (
                len(downloaded_records)
                >= args.max_results
            ):
                break

            if args.delay > 0:

                time.sleep(
                    args.delay
                )

        # Target reached.
        if (
            len(downloaded_records)
            >= args.max_results
        ):
            break

        # Get the next search page.
        if next_offset is not None:

            offset = next_offset

        else:

            offset += len(papers)

        # Fewer results than requested means we have
        # reached the end of the available results.
        if (
            len(papers)
            < args.page_size
        ):

            print()
            print(
                "Reached the end of the "
                "available search results."
            )

            break

        print()
        print(
            f"Waiting {args.search_delay}s "
            "before the next search page..."
        )

        if args.search_delay > 0:

            time.sleep(
                args.search_delay
            )

    # Always save the final Excel file.
    create_excel(
        excel_path,
        downloaded_records,
    )

    print()
    print("=" * 70)
    print("FINISHED")
    print("=" * 70)

    print(
        f"Downloaded PDFs: "
        f"{len(downloaded_records)}"
    )

    print(
        f"Download folder: "
        f"{output_dir}"
    )

    print(
        f"Excel file:      "
        f"{excel_path}"
    )

    print("=" * 70)

    if (
        len(downloaded_records)
        < args.max_results
    ):

        print()
        print(
            "WARNING:"
        )

        print(
            f"Requested {args.max_results} PDFs, "
            f"but only "
            f"{len(downloaded_records)} "
            "were successfully downloaded."
        )

        print(
            "The Excel file contains only "
            "successful downloads."
        )


if __name__ == "__main__":
    main()
